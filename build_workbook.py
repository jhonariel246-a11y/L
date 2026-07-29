#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Construye el libro resuelto del Ejercicio Practico Unidad 1 (PCT)."""
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side, NamedStyle
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import CellIsRule
from openpyxl.chart import BarChart, PieChart, Reference

SRC = "/root/.claude/uploads/db6384f2-a612-5778-9f45-9f62d3625d54/d0eb2a5a-Ejercicio_Practico_Unidad1_PCT_CON_INSTRUCCIONES.xlsx"
OUT = "/home/user/L/Ejercicio_Practico_Unidad1_PCT_RESUELTO.xlsx"

wb = openpyxl.load_workbook(SRC)
vt = wb["Ventas"]

# ---- Paleta ----
AZUL   = "1F4E78"
AZUL2  = "2E75B6"
GRIS   = "D9E1F2"
BLANCO = "FFFFFF"
FONT   = "Arial"

thin = Side(style="thin", color="BFBFBF")
borde = Border(left=thin, right=thin, top=thin, bottom=thin)

# ---- 1) Leer datos originales (encabezado fila 1, datos 2..51) ----
prods_cat = {
    "Impresora": "Tecnología", "Toner": "Tecnología", "Router": "Tecnología",
    "Tinta": "Tecnología", "Monitor": "Tecnología",
    "Carpeta": "Papelería", "Lápiz": "Papelería", "Cuadernos": "Papelería",
    "Escritorio": "Oficina", "Calculadora": "Oficina",
}
datos = []
for r in range(2, vt.max_row + 1):
    mes = vt.cell(r, 1).value
    prod = vt.cell(r, 2).value
    if mes is None and prod is None:
        continue
    cant = vt.cell(r, 4).value
    prec = vt.cell(r, 5).value
    if prod is None:
        continue
    cat = prods_cat.get(prod, "Oficina")
    datos.append((mes, prod, cat, cant, prec))

# ---- 2) Reconstruir hoja Ventas ----
# Borrar contenido previo
for row in vt.iter_rows():
    for c in row:
        c.value = None
        c.fill = PatternFill()
        c.border = Border()

# Encabezado / titulo
vt.merge_cells("A1:H1")
vt.merge_cells("A2:H2")
vt.merge_cells("A3:H3")
vt["A1"] = "LOS INTELECTUALES S.A."
vt["A2"] = "REPORTE DE VENTAS"
vt["A3"] = "PRIMER TRIMESTRE 2026"
vt["A1"].font = Font(name=FONT, size=18, bold=True, color=BLANCO)
vt["A2"].font = Font(name=FONT, size=13, bold=True, color=BLANCO)
vt["A3"].font = Font(name=FONT, size=12, bold=True, color="FFF2CC")
for cc in ("A1", "A2", "A3"):
    vt[cc].alignment = Alignment(horizontal="center", vertical="center")
    vt[cc].fill = PatternFill("solid", fgColor=AZUL)
vt.row_dimensions[1].height = 30
vt.row_dimensions[2].height = 22
vt.row_dimensions[3].height = 20

# Parametro IVA (referencia absoluta) -> celda K2
vt["J1"] = "PARÁMETROS"
vt["J1"].font = Font(name=FONT, size=10, bold=True, color=BLANCO)
vt["J1"].fill = PatternFill("solid", fgColor=AZUL2)
vt["J1"].alignment = Alignment(horizontal="center")
vt.merge_cells("J1:K1")
vt["J2"] = "IVA %"
vt["K2"] = 0.15
vt["J2"].font = Font(name=FONT, size=10, bold=True)
vt["K2"].font = Font(name=FONT, size=10, bold=True, color="C00000")
vt["K2"].number_format = "0%"
vt["J2"].fill = PatternFill("solid", fgColor="FFF2CC")
vt["K2"].fill = PatternFill("solid", fgColor="FFF2CC")
vt["J2"].border = borde
vt["K2"].border = borde
vt["J3"] = "(Ecuador 2026)"
vt["J3"].font = Font(name=FONT, size=8, italic=True, color="808080")
vt.merge_cells("J3:K3")

# Encabezados de tabla -> fila 5
headers = ["Mes", "Producto", "Categoría", "Cantidad\nVendida",
           "Precio\nUnitario", "Subtotal", "IVA", "Total\nVenta"]
HEAD_ROW = 5
for i, h in enumerate(headers, start=1):
    c = vt.cell(HEAD_ROW, i, h)
    c.font = Font(name=FONT, size=11, bold=True, color=BLANCO)
    c.fill = PatternFill("solid", fgColor=AZUL2)
    c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    c.border = borde
vt.row_dimensions[HEAD_ROW].height = 32

# Datos desde fila 6
FIRST = HEAD_ROW + 1
for idx, (mes, prod, cat, cant, prec) in enumerate(datos):
    r = FIRST + idx
    vt.cell(r, 1, mes)
    vt.cell(r, 2, prod)
    vt.cell(r, 3, cat)
    vt.cell(r, 4, cant)
    vt.cell(r, 5, prec)
    # Formulas
    vt.cell(r, 6, f"=D{r}*E{r}")           # Subtotal = Cantidad * Precio
    vt.cell(r, 7, f"=F{r}*$K$2")           # IVA con referencia absoluta
    vt.cell(r, 8, f"=F{r}+G{r}")           # Total Venta = Subtotal + IVA
LAST = FIRST + len(datos) - 1

# Fila de totales
TOT = LAST + 1
vt.cell(TOT, 5, "TOTALES:")
vt.cell(TOT, 5).font = Font(name=FONT, size=11, bold=True, color=BLANCO)
vt.cell(TOT, 5).alignment = Alignment(horizontal="right")
for col in (6, 7, 8):
    L = get_column_letter(col)
    c = vt.cell(TOT, col, f"=SUM({L}{FIRST}:{L}{LAST})")
    c.font = Font(name=FONT, size=11, bold=True, color=BLANCO)
    c.number_format = '"$"#,##0.00'
    c.fill = PatternFill("solid", fgColor=AZUL)
    c.border = borde
vt.cell(TOT, 5).fill = PatternFill("solid", fgColor=AZUL)
vt.cell(TOT, 5).border = borde

# Estilo de filas de datos
for r in range(FIRST, LAST + 1):
    zebra = GRIS if (r - FIRST) % 2 else BLANCO
    for col in range(1, 9):
        c = vt.cell(r, col)
        c.font = Font(name=FONT, size=10)
        c.border = borde
        c.fill = PatternFill("solid", fgColor=zebra)
        if col in (1, 2, 3):
            c.alignment = Alignment(horizontal="left", vertical="center")
        elif col == 4:
            c.alignment = Alignment(horizontal="center", vertical="center")
            c.number_format = "#,##0"
        else:
            c.alignment = Alignment(horizontal="right", vertical="center")
            c.number_format = '"$"#,##0.00'

# Anchos de columna
anchos = {"A": 12, "B": 14, "C": 13, "D": 11, "E": 11, "F": 12, "G": 11, "H": 13,
          "I": 3, "J": 12, "K": 10}
for col, w in anchos.items():
    vt.column_dimensions[col].width = w

# ---- 3) Formato condicional en Total Venta (col H) ----
rango_h = f"H{FIRST}:H{LAST}"
verde = PatternFill("solid", fgColor="C6EFCE")
rojo  = PatternFill("solid", fgColor="FFC7CE")
vt.conditional_formatting.add(rango_h,
    CellIsRule(operator="greaterThan", formula=["500"], fill=verde,
               font=Font(name=FONT, color="006100")))
vt.conditional_formatting.add(rango_h,
    CellIsRule(operator="lessThan", formula=["200"], fill=rojo,
               font=Font(name=FONT, color="9C0006")))

# ---- 4) Lista desplegable Categoria (Hoja1!A3:A5) ----
dv = DataValidation(type="list", formula1="=Hoja1!$A$3:$A$5", allow_blank=False,
                    showDropDown=False)
dv.error = "Seleccione: Papelería, Tecnología u Oficina"
dv.errorTitle = "Categoría inválida"
dv.prompt = "Elija una categoría de la lista"
dv.promptTitle = "Categoría"
vt.add_data_validation(dv)
dv.add(f"C{FIRST}:C{LAST}")

vt.freeze_panes = "A6"
vt.sheet_view.showGridLines = False

# Mejorar hoja Hoja1 (lista de categorias)
h1 = wb["Hoja1"]
h1["A1"] = "CATEGORÍAS"
h1["A1"].font = Font(name=FONT, size=11, bold=True, color=BLANCO)
h1["A1"].fill = PatternFill("solid", fgColor=AZUL2)
h1["A1"].alignment = Alignment(horizontal="center")
h1.merge_cells("A1:A2")
for rr in (3, 4, 5):
    c = h1.cell(rr, 1)
    c.font = Font(name=FONT, size=10)
    c.border = borde
    c.fill = PatternFill("solid", fgColor=GRIS if rr % 2 else BLANCO)
h1.column_dimensions["A"].width = 16

# ---- 5-7) Hoja Resumen + graficos + tabla tipo dinamica ----
if "Resumen y Gráficos" in wb.sheetnames:
    del wb["Resumen y Gráficos"]
rs = wb.create_sheet("Resumen y Gráficos")
rs.sheet_view.showGridLines = False

rs["A1"] = "LOS INTELECTUALES S.A.  —  ANÁLISIS DE VENTAS 1er TRIMESTRE 2026"
rs.merge_cells("A1:H1")
rs["A1"].font = Font(name=FONT, size=14, bold=True, color=BLANCO)
rs["A1"].fill = PatternFill("solid", fgColor=AZUL)
rs["A1"].alignment = Alignment(horizontal="center", vertical="center")
rs.row_dimensions[1].height = 26

def titulo(celda, texto):
    rs[celda] = texto
    rs[celda].font = Font(name=FONT, size=11, bold=True, color=BLANCO)
    rs[celda].fill = PatternFill("solid", fgColor=AZUL2)
    rs[celda].alignment = Alignment(horizontal="center")

def encab(cell, texto):
    rs[cell] = texto
    rs[cell].font = Font(name=FONT, size=10, bold=True, color=BLANCO)
    rs[cell].fill = PatternFill("solid", fgColor=AZUL2)
    rs[cell].alignment = Alignment(horizontal="center")
    rs[cell].border = borde

vrango_mes = f"Ventas!$A${FIRST}:$A${LAST}"
vrango_tot = f"Ventas!$H${FIRST}:$H${LAST}"
vrango_cat = f"Ventas!$C${FIRST}:$C${LAST}"

# Tabla: Ventas por Mes
titulo("A3", "VENTAS TOTALES POR MES")
rs.merge_cells("A3:B3")
encab("A4", "Mes"); encab("B4", "Total Venta")
meses = ["Enero", "Febrero", "Marzo"]
for i, m in enumerate(meses):
    r = 5 + i
    rs.cell(r, 1, m).border = borde
    rs.cell(r, 1).font = Font(name=FONT, size=10)
    c = rs.cell(r, 2, f'=SUMIFS({vrango_tot},{vrango_mes},A{r})')
    c.number_format = '"$"#,##0.00'; c.border = borde
    c.font = Font(name=FONT, size=10)
rs.cell(8, 1, "TOTAL").font = Font(name=FONT, size=10, bold=True)
rs.cell(8, 1).border = borde
ct = rs.cell(8, 2, "=SUM(B5:B7)")
ct.number_format = '"$"#,##0.00'; ct.border = borde
ct.font = Font(name=FONT, size=10, bold=True)

# Tabla: Ventas por Categoria
titulo("D3", "VENTAS POR CATEGORÍA")
rs.merge_cells("D3:E3")
encab("D4", "Categoría"); encab("E4", "Total Venta")
cats = ["Tecnología", "Papelería", "Oficina"]
for i, cat in enumerate(cats):
    r = 5 + i
    rs.cell(r, 4, cat).border = borde
    rs.cell(r, 4).font = Font(name=FONT, size=10)
    c = rs.cell(r, 5, f'=SUMIFS({vrango_tot},{vrango_cat},D{r})')
    c.number_format = '"$"#,##0.00'; c.border = borde
    c.font = Font(name=FONT, size=10)
rs.cell(8, 4, "TOTAL").font = Font(name=FONT, size=10, bold=True)
rs.cell(8, 4).border = borde
cc = rs.cell(8, 5, "=SUM(E5:E7)")
cc.number_format = '"$"#,##0.00'; cc.border = borde
cc.font = Font(name=FONT, size=10, bold=True)

# Tabla tipo DINAMICA: Categoria x Mes
titulo("A11", "TABLA DINÁMICA — VENTAS POR CATEGORÍA Y MES")
rs.merge_cells("A11:E11")
encab("A12", "Categoría \\ Mes")
for j, m in enumerate(meses):
    encab(get_column_letter(2 + j) + "12", m)
encab("E12", "Total general")
for i, cat in enumerate(cats):
    r = 13 + i
    rs.cell(r, 1, cat).border = borde
    rs.cell(r, 1).font = Font(name=FONT, size=10, bold=True)
    for j, m in enumerate(meses):
        col = 2 + j
        L = get_column_letter(col)
        c = rs.cell(r, col,
            f'=SUMIFS({vrango_tot},{vrango_cat},$A{r},{vrango_mes},{L}$12)')
        c.number_format = '"$"#,##0.00'; c.border = borde
        c.font = Font(name=FONT, size=10)
    tc = rs.cell(r, 5, f"=SUM(B{r}:D{r})")
    tc.number_format = '"$"#,##0.00'; tc.border = borde
    tc.font = Font(name=FONT, size=10, bold=True)
# Fila total general
rtot = 13 + len(cats)
rs.cell(rtot, 1, "Total general").font = Font(name=FONT, size=10, bold=True)
rs.cell(rtot, 1).border = borde
for col in range(2, 6):
    L = get_column_letter(col)
    c = rs.cell(rtot, col, f"=SUM({L}13:{L}{rtot-1})")
    c.number_format = '"$"#,##0.00'; c.border = borde
    c.font = Font(name=FONT, size=10, bold=True)
    c.fill = PatternFill("solid", fgColor=GRIS)
rs.cell(rtot, 1).fill = PatternFill("solid", fgColor=GRIS)

nota = rs.cell(rtot + 2, 1,
    "Nota: Tabla de resumen (SUMIFS) equivalente a una tabla dinámica de "
    "Categoría × Mes. En Excel puede convertirla en Tabla Dinámica nativa "
    "(Insertar > Tabla dinámica) y añadir un Segmentador por Producto.")
nota.font = Font(name=FONT, size=8, italic=True, color="808080")
rs.merge_cells(f"A{rtot+2}:H{rtot+3}")
nota.alignment = Alignment(wrap_text=True, vertical="top")

for col, w in {"A": 20, "B": 13, "C": 13, "D": 14, "E": 14}.items():
    rs.column_dimensions[col].width = w

# ---- Grafico de columnas: ventas por mes ----
bar = BarChart()
bar.type = "col"
bar.title = "Ventas Totales por Mes"
bar.style = 10
bar.y_axis.title = "Total Venta ($)"
bar.x_axis.title = "Mes"
data_ref = Reference(rs, min_col=2, min_row=4, max_row=7)
cats_ref = Reference(rs, min_col=1, min_row=5, max_row=7)
bar.add_data(data_ref, titles_from_data=True)
bar.set_categories(cats_ref)
bar.height = 8; bar.width = 15
bar.legend = None
rs.add_chart(bar, "A21")

# ---- Grafico circular: ventas por categoria ----
pie = PieChart()
pie.title = "Ventas por Categoría"
pie.style = 10
pdata = Reference(rs, min_col=5, min_row=4, max_row=7)
plabels = Reference(rs, min_col=4, min_row=5, max_row=7)
pie.add_data(pdata, titles_from_data=True)
pie.set_categories(plabels)
pie.height = 8; pie.width = 12
from openpyxl.chart.label import DataLabelList
pie.dataLabels = DataLabelList()
pie.dataLabels.showPercent = True
rs.add_chart(pie, "J21")

# ---- Ordenar hojas ----
wb.move_sheet("Resumen y Gráficos", -(wb.sheetnames.index("Resumen y Gráficos") - 2))

wb.save(OUT)
print("Guardado:", OUT)
print("Filas de datos:", len(datos), "-> filas", FIRST, "a", LAST, "| Totales fila", TOT)
