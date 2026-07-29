#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Inyecta valores en cache <v> en las celdas con formula (sin LibreOffice)."""
import zipfile, shutil, re, os
from openpyxl import load_workbook

F = "/home/user/L/Ejercicio_Practico_Unidad1_PCT_RESUELTO.xlsx"

wb = load_workbook(F)
vt = wb["Ventas"]
rs = wb["Resumen y Gráficos"]

IVA = 0.15
FIRST, LAST = 6, 55

# --- Calcular valores Ventas ---
vals = {}  # (sheetname, coord) -> number
sub_by_row = {}
for r in range(FIRST, LAST + 1):
    d = vt.cell(r, 4).value
    e = vt.cell(r, 5).value
    sub = d * e
    iva = sub * IVA
    tot = sub + iva
    sub_by_row[r] = (sub, iva, tot)
    vals[("Ventas", f"F{r}")] = sub
    vals[("Ventas", f"G{r}")] = iva
    vals[("Ventas", f"H{r}")] = tot
vals[("Ventas", "F56")] = sum(v[0] for v in sub_by_row.values())
vals[("Ventas", "G56")] = sum(v[1] for v in sub_by_row.values())
vals[("Ventas", "H56")] = sum(v[2] for v in sub_by_row.values())

# --- Resumen: por mes / categoria / cross-tab ---
rows = []
for r in range(FIRST, LAST + 1):
    rows.append((vt.cell(r, 1).value, vt.cell(r, 3).value, sub_by_row[r][2]))

meses = ["Enero", "Febrero", "Marzo"]
cats = ["Tecnología", "Papelería", "Oficina"]

RS = "Resumen y Gráficos"
# por mes B5:B7, B8
for i, m in enumerate(meses):
    tot = sum(t for mm, cc, t in rows if mm == m)
    vals[(RS, f"B{5+i}")] = tot
vals[(RS, "B8")] = sum(vals[(RS, f"B{5+i}")] for i in range(3))
# por categoria E5:E7, E8
for i, c in enumerate(cats):
    tot = sum(t for mm, cc, t in rows if cc == c)
    vals[(RS, f"E{5+i}")] = tot
vals[(RS, "E8")] = sum(vals[(RS, f"E{5+i}")] for i in range(3))
# cross-tab B13:D15 (cats rows, months cols), E row totals
colL = {0: "B", 1: "C", 2: "D"}
for ci, c in enumerate(cats):
    r = 13 + ci
    rowtot = 0
    for mi, m in enumerate(meses):
        v = sum(t for mm, cc, t in rows if cc == c and mm == m)
        vals[(RS, f"{colL[mi]}{r}")] = v
        rowtot += v
    vals[(RS, f"E{r}")] = rowtot
# total general row 16
for L in ("B", "C", "D", "E"):
    vals[(RS, f"{L}16")] = sum(vals[(RS, f"{L}{13+ci}")] for ci in range(3))

# --- Mapear nombre de hoja -> archivo xml ---
tmp = F + ".tmp.zip"
with zipfile.ZipFile(F) as z:
    wbxml = z.read("xl/workbook.xml").decode("utf-8")
    relsxml = z.read("xl/_rels/workbook.xml.rels").decode("utf-8")
    names = z.namelist()

# sheet name -> r:id
sheet_rid = dict(re.findall(r'<sheet[^>]*name="([^"]*)"[^>]*r:id="([^"]*)"', wbxml))
# también orden alternativo de atributos
if not sheet_rid:
    for m in re.finditer(r'<sheet\b[^>]*/>', wbxml):
        tag = m.group(0)
        nm = re.search(r'name="([^"]*)"', tag)
        rid = re.search(r'r:id="([^"]*)"', tag)
        if nm and rid:
            sheet_rid[nm.group(1)] = rid.group(1)
# rid -> target (los atributos pueden venir en cualquier orden)
rid_target = {}
for tag in re.findall(r'<Relationship\b[^>]*/>', relsxml):
    rid = re.search(r'Id="([^"]*)"', tag)
    tgt = re.search(r'Target="([^"]*)"', tag)
    if rid and tgt:
        rid_target[rid.group(1)] = tgt.group(1)
sheet_file = {}
for nm, rid in sheet_rid.items():
    t = rid_target[rid].lstrip("/")
    if not t.startswith("xl/"):
        t = "xl/" + t
    sheet_file[nm] = t

def fmt(v):
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return repr(float(v))

# Agrupar valores por archivo
by_file = {}
for (sheet, coord), v in vals.items():
    by_file.setdefault(sheet_file[sheet], {})[coord] = v

# --- Reescribir el zip con <v> inyectado ---
with zipfile.ZipFile(F) as zin, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
    for item in zin.infolist():
        data = zin.read(item.filename)
        if item.filename in by_file:
            txt = data.decode("utf-8")
            for coord, v in by_file[item.filename].items():
                # celda con formula: <c r="F6" ...><f>...</f></c>  -> añadir <v>
                pat = re.compile(r'(<c r="' + coord + r'"[^>]*>)(.*?)(</c>)', re.S)
                def repl(m, v=v):
                    inner = m.group(2)
                    inner = re.sub(r'<v\s*/>', '', inner)
                    inner = re.sub(r'<v>.*?</v>', '', inner, flags=re.S)
                    return m.group(1) + inner + f'<v>{fmt(v)}</v>' + m.group(3)
                txt, n = pat.subn(repl, txt)
            data = txt.encode("utf-8")
        zout.writestr(item, data)

shutil.move(tmp, F)
print("Valores inyectados:", len(vals))
# Verificar
wb2 = load_workbook(F, data_only=True)
print("Check Ventas H6 =", wb2["Ventas"]["H6"].value, "(esperado 862.5)")
print("Check Ventas H56 (total) =", wb2["Ventas"]["H56"].value)
print("Check Resumen B5 (Enero) =", wb2["Resumen y Gráficos"]["B5"].value)
print("Check Resumen E8 (total cat) =", wb2["Resumen y Gráficos"]["E8"].value)
