# -*- coding: utf-8 -*-
"""Inyecta una TABLA DINÁMICA real + SEGMENTADOR (slicer) en el .xlsx.

Construye a mano las partes OOXML (pivotCache, pivotTable, slicerCache,
slicer, drawing) y las conecta en workbook/worksheet/content-types/rels.
Se usa refreshOnLoad para que Excel reconstruya el cuerpo de la dinámica
al abrir, evitando depender de un cuerpo pre-renderizado exacto.
"""
import os
import zipfile
import openpyxl

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
XLSX = os.path.join(ROOT, 'Reporte_Ventas_Jhon_Martinez.xlsx')
SRC = os.path.join(ROOT, 'archivo_original.xlsx')

CATEGORIA = {
    'Impresora': 'Tecnología', 'Toner': 'Tecnología', 'Router': 'Tecnología',
    'Tinta': 'Tecnología', 'Monitor': 'Tecnología', 'Calculadora': 'Tecnología',
    'Carpeta': 'Papelería', 'Lápiz': 'Papelería', 'Cuadernos': 'Papelería',
    'Escritorio': 'Oficina',
}
MESES = ['Enero', 'Febrero', 'Marzo']
CATS = ['Tecnología', 'Papelería', 'Oficina']
IVA = 0.15


def esc(s):
    return (s.replace('&', '&amp;').replace('<', '&lt;').replace('>', '&gt;')
             .replace('"', '&quot;'))


def leer_registros():
    wb = openpyxl.load_workbook(SRC)
    ws = wb['Ventas']
    regs = []
    for r in range(2, 52):
        mes = ws.cell(r, 1).value
        prod = ws.cell(r, 2).value
        cant = ws.cell(r, 4).value
        precio = ws.cell(r, 5).value
        if mes and prod:
            regs.append((mes, prod, cant, precio))
    return regs


def fnum(x):
    if float(x).is_integer():
        return str(int(x))
    return repr(round(float(x), 6))


def main():
    regs = leer_registros()
    productos = []
    for _, p, _, _ in regs:
        if p not in productos:
            productos.append(p)

    mes_idx = {m: i for i, m in enumerate(MESES)}
    prod_idx = {p: i for i, p in enumerate(productos)}
    cat_idx = {c: i for i, c in enumerate(CATS)}

    # ---------- pivotCacheRecords1.xml ----------
    recs = []
    cants, precios, subs, ivas, tots = [], [], [], [], []
    for mes, prod, cant, precio in regs:
        cat = CATEGORIA[prod]
        sub = cant * precio
        iva = sub * IVA
        tot = sub + iva
        cants.append(cant); precios.append(precio)
        subs.append(sub); ivas.append(iva); tots.append(tot)
        recs.append(
            '<r><x v="%d"/><x v="%d"/><x v="%d"/><n v="%s"/><n v="%s"/>'
            '<n v="%s"/><n v="%s"/><n v="%s"/></r>' % (
                mes_idx[mes], prod_idx[prod], cat_idx[cat],
                fnum(cant), fnum(precio), fnum(sub), fnum(iva), fnum(tot)))
    records = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
        '<pivotCacheRecords xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" '
        'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" '
        'count="%d">%s</pivotCacheRecords>' % (len(recs), ''.join(recs)))

    # ---------- pivotCacheDefinition1.xml ----------
    def shared_str(vals):
        return ''.join('<s v="%s"/>' % esc(v) for v in vals)

    def num_field(name, vals, integer=False):
        mn = min(vals); mx = max(vals)
        extra = ' containsInteger="1"' if integer else ''
        return ('<cacheField name="%s" numFmtId="0"><sharedItems '
                'containsSemiMixedTypes="0" containsString="0" '
                'containsNumber="1"%s minValue="%s" maxValue="%s"/></cacheField>'
                % (esc(name), extra, fnum(mn), fnum(mx)))

    cache_fields = (
        '<cacheField name="Mes" numFmtId="0"><sharedItems count="%d">%s</sharedItems></cacheField>'
        % (len(MESES), shared_str(MESES)) +
        '<cacheField name="Producto" numFmtId="0"><sharedItems count="%d">%s</sharedItems></cacheField>'
        % (len(productos), shared_str(productos)) +
        '<cacheField name="Categoría" numFmtId="0"><sharedItems count="%d">%s</sharedItems></cacheField>'
        % (len(CATS), shared_str(CATS)) +
        num_field('Cantidad Vendida', cants, integer=all(float(v).is_integer() for v in cants)) +
        num_field('Precio Unitario', precios, integer=all(float(v).is_integer() for v in precios)) +
        num_field('Subtotal', subs) +
        num_field('IVA', ivas) +
        num_field('Total Venta', tots))

    cache_def = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
        '<pivotCacheDefinition xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" '
        'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" '
        'r:id="rId1" refreshedBy="Excel" refreshedDate="46000" createdVersion="6" '
        'refreshedVersion="6" minRefreshableVersion="3" recordCount="%d" '
        'refreshOnLoad="1" saveData="1"><cacheSource type="worksheet">'
        '<worksheetSource ref="A5:H55" sheet="Ventas"/></cacheSource>'
        '<cacheFields count="8">%s</cacheFields></pivotCacheDefinition>'
        % (len(recs), cache_fields))

    # ---------- pivotTable1.xml ----------
    def axis_items(n):
        it = ''.join('<item x="%d"/>' % i for i in range(n))
        return '<items count="%d">%s<item t="default"/></items>' % (n + 1, it)

    pivot_fields = (
        '<pivotField axis="axisCol" showAll="0">%s</pivotField>' % axis_items(len(MESES)) +   # Mes
        '<pivotField showAll="0"/>' +                                                          # Producto
        '<pivotField axis="axisRow" showAll="0">%s</pivotField>' % axis_items(len(CATS)) +     # Categoría
        '<pivotField showAll="0"/>' +   # Cantidad
        '<pivotField showAll="0"/>' +   # Precio
        '<pivotField showAll="0"/>' +   # Subtotal
        '<pivotField showAll="0"/>' +   # IVA
        '<pivotField dataField="1" showAll="0"/>')  # Total Venta

    def line_items(n):
        parts = ['<i><x/></i>'] + ['<i><x v="%d"/></i>' % i for i in range(1, n)]
        parts.append('<i t="grand"><x/></i>')
        return ''.join(parts)

    pivot_table = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
        '<pivotTableDefinition xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" '
        'name="TablaDinamica1" cacheId="1" applyNumberFormats="0" applyBorderFormats="0" '
        'applyFontFormats="0" applyPatternFormats="0" applyAlignmentFormats="0" '
        'applyWidthHeightFormats="1" dataCaption="Valores" updatedVersion="6" '
        'minRefreshableVersion="3" useAutoFormatting="1" itemPrintTitles="1" '
        'createdVersion="6" indent="0" outline="1" outlineData="1" multipleFieldFilters="0">'
        '<location ref="A4:E9" firstHeaderRow="1" firstDataRow="2" firstDataCol="1"/>'
        '<pivotFields count="8">%s</pivotFields>'
        '<rowFields count="1"><field x="2"/></rowFields>'
        '<rowItems count="%d">%s</rowItems>'
        '<colFields count="1"><field x="0"/></colFields>'
        '<colItems count="%d">%s</colItems>'
        '<dataFields count="1"><dataField name="Suma de Total Venta" fld="7" baseField="0" baseItem="0" numFmtId="164"/></dataFields>'
        '<pivotTableStyleInfo name="PivotStyleLight16" showRowHeaders="1" showColHeaders="1" '
        'showRowStripes="0" showColStripes="0" showLastColumn="1"/>'
        '</pivotTableDefinition>' % (
            pivot_fields, len(CATS) + 1, line_items(len(CATS)),
            len(MESES) + 1, line_items(len(MESES))))

    # ---------- slicerCache1.xml ----------
    slicer_items = ''.join('<i x="%d"/>' % i for i in range(len(productos)))
    slicer_cache = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
        '<slicerCacheDefinition xmlns="http://schemas.microsoft.com/office/spreadsheetml/2009/9/main" '
        'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" '
        'name="Slicer_Producto" sourceName="Producto">'
        '<pivotTables><pivotTable tabId="5" name="TablaDinamica1"/></pivotTables>'
        '<data><tabular pivotCacheId="1"><items count="%d">%s</items></tabular></data>'
        '</slicerCacheDefinition>' % (len(productos), slicer_items))

    # ---------- slicer1.xml ----------
    slicer = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
        '<slicers xmlns="http://schemas.microsoft.com/office/spreadsheetml/2009/9/main" '
        'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" '
        'xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006" '
        'xmlns:x="http://schemas.openxmlformats.org/spreadsheetml/2006/main">'
        '<slicer name="Producto" cache="Slicer_Producto" caption="Producto" '
        'startItem="0" columnCount="1" showCaption="1" level="0" style="SlicerStyleLight1" rowHeight="241300"/>'
        '</slicers>')

    # ---------- drawing2.xml (forma visual del slicer) ----------
    drawing = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
        '<xdr:wsDr xmlns:xdr="http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing" '
        'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">'
        '<mc:AlternateContent xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006">'
        '<mc:Choice xmlns:sle="http://schemas.microsoft.com/office/drawing/2010/slicer" Requires="sle">'
        '<xdr:twoCellAnchor>'
        '<xdr:from><xdr:col>6</xdr:col><xdr:colOff>0</xdr:colOff><xdr:row>3</xdr:row><xdr:rowOff>0</xdr:rowOff></xdr:from>'
        '<xdr:to><xdr:col>8</xdr:col><xdr:colOff>0</xdr:colOff><xdr:row>18</xdr:row><xdr:rowOff>0</xdr:rowOff></xdr:to>'
        '<xdr:graphicFrame macro="">'
        '<xdr:nvGraphicFramePr><xdr:cNvPr id="2" name="Producto"/><xdr:cNvGraphicFramePr/></xdr:nvGraphicFramePr>'
        '<xdr:xfrm><a:off x="0" y="0"/><a:ext cx="0" cy="0"/></xdr:xfrm>'
        '<a:graphic><a:graphicData uri="http://schemas.microsoft.com/office/drawing/2010/slicer">'
        '<sle:slicer xmlns:sle="http://schemas.microsoft.com/office/drawing/2010/slicer" name="Producto"/>'
        '</a:graphicData></a:graphic></xdr:graphicFrame>'
        '<xdr:clientData/></xdr:twoCellAnchor></mc:Choice>'
        '<mc:Fallback><xdr:twoCellAnchor>'
        '<xdr:from><xdr:col>6</xdr:col><xdr:colOff>0</xdr:colOff><xdr:row>3</xdr:row><xdr:rowOff>0</xdr:rowOff></xdr:from>'
        '<xdr:to><xdr:col>8</xdr:col><xdr:colOff>0</xdr:colOff><xdr:row>18</xdr:row><xdr:rowOff>0</xdr:rowOff></xdr:to>'
        '<xdr:sp macro="" textlink=""><xdr:nvSpPr><xdr:cNvPr id="2" name="Producto"/>'
        '<xdr:cNvSpPr><a:spLocks noTextEdit="1"/></xdr:cNvSpPr></xdr:nvSpPr>'
        '<xdr:spPr><a:xfrm><a:off x="0" y="0"/><a:ext cx="0" cy="0"/></a:xfrm>'
        '<a:prstGeom prst="rect"><a:avLst/></a:prstGeom></xdr:spPr>'
        '<xdr:txBody><a:bodyPr vertOverflow="clip" horzOverflow="clip"/><a:lstStyle/>'
        '<a:p><a:r><a:rPr lang="es-EC"/><a:t>Segmentador: Producto</a:t></a:r></a:p></xdr:txBody>'
        '</xdr:sp><xdr:clientData/></xdr:twoCellAnchor></mc:Fallback>'
        '</mc:AlternateContent></xdr:wsDr>')

    # ---------- relationships ----------
    cache_def_rels = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
        '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
        '<Relationship Id="rId1" '
        'Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/pivotCacheRecords" '
        'Target="pivotCacheRecords1.xml"/></Relationships>')

    pivot_table_rels = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
        '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
        '<Relationship Id="rId1" '
        'Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/pivotCacheDefinition" '
        'Target="../pivotCache/pivotCacheDefinition1.xml"/></Relationships>')

    sheet5_rels = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
        '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
        '<Relationship Id="rId1" '
        'Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/drawing" '
        'Target="../drawings/drawing2.xml"/>'
        '<Relationship Id="rId2" '
        'Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/pivotTable" '
        'Target="../pivotTables/pivotTable1.xml"/>'
        '<Relationship Id="rId3" '
        'Type="http://schemas.microsoft.com/office/2007/relationships/slicer" '
        'Target="../slicers/slicer1.xml"/></Relationships>')

    # ---------- read existing zip ----------
    zin = zipfile.ZipFile(XLSX, 'r')
    parts = {n: zin.read(n) for n in zin.namelist()}
    zin.close()

    # patch [Content_Types].xml
    ct = parts['[Content_Types].xml'].decode('utf-8')
    overrides = (
        '<Override PartName="/xl/pivotCache/pivotCacheDefinition1.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.pivotCacheDefinition+xml"/>'
        '<Override PartName="/xl/pivotCache/pivotCacheRecords1.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.pivotCacheRecords+xml"/>'
        '<Override PartName="/xl/pivotTables/pivotTable1.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.pivotTable+xml"/>'
        '<Override PartName="/xl/slicerCaches/slicerCache1.xml" ContentType="application/vnd.ms-excel.slicerCache+xml"/>'
        '<Override PartName="/xl/slicers/slicer1.xml" ContentType="application/vnd.ms-excel.slicer+xml"/>'
        '<Override PartName="/xl/drawings/drawing2.xml" ContentType="application/vnd.openxmlformats-officedocument.drawing+xml"/>')
    ct = ct.replace('</Types>', overrides + '</Types>')
    parts['[Content_Types].xml'] = ct.encode('utf-8')

    # patch workbook.xml : pivotCaches (after calcPr) + extLst (slicerCaches)
    wbx = parts['xl/workbook.xml'].decode('utf-8')
    pivot_caches = '<pivotCaches><pivotCache cacheId="1" r:id="rId8"/></pivotCaches>'
    wb_ext = ('<extLst><ext xmlns:x14="http://schemas.microsoft.com/office/spreadsheetml/2009/9/main" '
              'uri="{BBE1A952-AA13-448e-AADC-164F8A28A991}"><x14:slicerCaches>'
              '<x14:slicerCache r:id="rId9"/></x14:slicerCaches></ext></extLst>')
    if '</calcPr>' in wbx:
        wbx = wbx.replace('</calcPr>', '</calcPr>' + pivot_caches, 1)
    elif '<calcPr' in wbx:
        import re
        wbx = re.sub(r'(<calcPr[^>]*/>)', r'\1' + pivot_caches, wbx, count=1)
    else:
        wbx = wbx.replace('</sheets>', '</sheets>' + pivot_caches, 1)
    wbx = wbx.replace('</workbook>', wb_ext + '</workbook>')
    parts['xl/workbook.xml'] = wbx.encode('utf-8')

    # patch workbook rels
    wbr = parts['xl/_rels/workbook.xml.rels'].decode('utf-8')
    add_rels = (
        '<Relationship Id="rId8" '
        'Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/pivotCacheDefinition" '
        'Target="pivotCache/pivotCacheDefinition1.xml"/>'
        '<Relationship Id="rId9" '
        'Type="http://schemas.microsoft.com/office/2007/relationships/slicerCache" '
        'Target="slicerCaches/slicerCache1.xml"/>')
    wbr = wbr.replace('</Relationships>', add_rels + '</Relationships>')
    parts['xl/_rels/workbook.xml.rels'] = wbr.encode('utf-8')

    # patch sheet5.xml : xmlns:r on root + drawing + extLst(slicerList)
    s5 = parts['xl/worksheets/sheet5.xml'].decode('utf-8')
    if 'xmlns:r=' not in s5.split('>', 1)[0]:
        s5 = s5.replace(
            '<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"',
            '<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" '
            'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"', 1)
    sheet_tail = ('<drawing r:id="rId1"/>'
                  '<extLst><ext xmlns:x14="http://schemas.microsoft.com/office/spreadsheetml/2009/9/main" '
                  'uri="{A8765BA9-456A-4dab-B4F3-ACF838C121DE}"><x14:slicerList>'
                  '<x14:slicer r:id="rId3"/></x14:slicerList></ext></extLst>')
    s5 = s5.replace('</worksheet>', sheet_tail + '</worksheet>')
    parts['xl/worksheets/sheet5.xml'] = s5.encode('utf-8')

    # ---------- write new parts ----------
    new = {
        'xl/pivotCache/pivotCacheDefinition1.xml': cache_def,
        'xl/pivotCache/pivotCacheRecords1.xml': records,
        'xl/pivotCache/_rels/pivotCacheDefinition1.xml.rels': cache_def_rels,
        'xl/pivotTables/pivotTable1.xml': pivot_table,
        'xl/pivotTables/_rels/pivotTable1.xml.rels': pivot_table_rels,
        'xl/slicerCaches/slicerCache1.xml': slicer_cache,
        'xl/slicers/slicer1.xml': slicer,
        'xl/drawings/drawing2.xml': drawing,
        'xl/worksheets/_rels/sheet5.xml.rels': sheet5_rels,
    }
    for k, v in new.items():
        parts[k] = v.encode('utf-8')

    # ---------- rewrite zip ----------
    tmp = XLSX + '.tmp'
    zout = zipfile.ZipFile(tmp, 'w', zipfile.ZIP_DEFLATED)
    for name, content in parts.items():
        zout.writestr(name, content)
    zout.close()
    os.replace(tmp, XLSX)
    print('Tabla dinámica + segmentador inyectados en', os.path.basename(XLSX))
    print('Productos (slicer):', productos)


if __name__ == '__main__':
    main()
